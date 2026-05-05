"""
Gera o arquivo `templates/obra-template.xlsx` programaticamente.

Este script é utilitário — roda **uma vez** para criar o template, e depois o
arquivo gerado é distribuído como knowledge file do Gem ou para download direto.

Uso (de dentro de ponto-de-bloqueio/):
    python -m scripts.generate_template
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# Paleta sóbria (paleta BravoCP futura: azul-marinho + cinza)
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
EXAMPLE_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
EXAMPLE_FONT = Font(italic=True, color="475569")

THIN = Side(border_style="thin", color="CBD5E1")
BORDER = Border(top=THIN, left=THIN, right=THIN, bottom=THIN)


COLUMNS = [
    ("ID", 12, "Auto-gerado pela skill (sigla do contratado + nº). Ex: TRI-01"),
    ("Contratado", 18, "Nome do fornecedor"),
    ("Escopo", 35, "Tarefa que esse contratado vai executar nessa frente"),
    ("Entra após", 18, "IDs separados por vírgula. Vazio = tarefa raiz."),
    ("Entrada planejada", 16, "Data prevista de entrada do contratado"),
    ("Saída planejada", 16, "Data prevista de saída do contratado"),
    ("Entrada real", 14, "Preencher quando o contratado efetivamente entrar"),
    ("Saída real", 14, "Preencher quando o contratado efetivamente sair"),
]


def _setup_obra_sheet(ws: Worksheet) -> None:
    """Aba principal — onde o usuário cadastra as tarefas."""
    ws.title = "Obra"

    # Linha 1: cabeçalhos
    for col_idx, (name, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Linha 2: exemplo
    example = [
        "TRI-01", "Triade", "Estrutura metálica Galpão 2", "",
        date(2026, 4, 1), date(2026, 4, 30), None, None,
    ]
    for col_idx, value in enumerate(example, start=1):
        cell = ws.cell(row=2, column=col_idx, value=value)
        cell.font = EXAMPLE_FONT
        cell.fill = EXAMPLE_FILL
        cell.border = BORDER
        if isinstance(value, date):
            cell.number_format = "DD/MM/YYYY"

    # Linha 3: segunda linha de exemplo (com dependência)
    example2 = [
        "SAU-01", "Saur", "Cobertura Galpão 2", "TRI-01",
        date(2026, 5, 1), date(2026, 5, 30), None, None,
    ]
    for col_idx, value in enumerate(example2, start=1):
        cell = ws.cell(row=3, column=col_idx, value=value)
        cell.font = EXAMPLE_FONT
        cell.fill = EXAMPLE_FILL
        cell.border = BORDER
        if isinstance(value, date):
            cell.number_format = "DD/MM/YYYY"

    # Apenas as 2 linhas de exemplo + cabeçalho. Bordas para 50 linhas extras (vazias)
    for row in range(4, 54):
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            if col >= 5:  # colunas de data
                cell.number_format = "DD/MM/YYYY"

    # Congelar cabeçalho
    ws.freeze_panes = "A2"

    # Altura da linha de cabeçalho
    ws.row_dimensions[1].height = 24


def _setup_como_usar_sheet(wb: Workbook) -> None:
    """Aba com instruções de uso."""
    ws = wb.create_sheet("Como usar")
    ws.column_dimensions["A"].width = 110

    title_font = Font(bold=True, size=14, color="1F2937")
    section_font = Font(bold=True, size=11, color="1F2937")
    body_font = Font(size=11, color="333333")

    content = [
        ("Como usar este template", title_font),
        ("", body_font),
        ("Esta planilha é a fonte da verdade da sua obra. "
         "Você a preenche, faz upload no Ponto de Bloqueio, e a IA devolve a análise.", body_font),
        ("", body_font),
        ("PASSO A PASSO", section_font),
        ("1. Liste cada tarefa em uma linha. Uma 'tarefa' = um contratado executando "
         "um escopo específico em uma frente da obra.", body_font),
        ("2. Não preencha a coluna ID — a IA gera para você no formato SIGLA-NN.", body_font),
        ("3. Em 'Entra após', coloque os IDs das tarefas que precisam terminar antes "
         "desta começar (separados por vírgula). Deixe vazio se for tarefa raiz.", body_font),
        ("4. Datas planejadas são obrigatórias. Datas reais (entrada/saída) só preencha "
         "quando o contratado efetivamente entrar ou sair do canteiro.", body_font),
        ("5. À medida que a obra avança, atualize as datas reais e suba a planilha "
         "novamente. A IA recalcula tudo.", body_font),
        ("", body_font),
        ("REGRAS IMPORTANTES", section_font),
        ("• Mesmo contratado pode aparecer em múltiplas linhas — uma por tarefa/frente.", body_font),
        ("• Cada tarefa tem APENAS dois marcos: quando entra e quando sai. "
         "Você NÃO precisa modelar o cronograma interno do fornecedor.", body_font),
        ("• Uma tarefa não pode depender, direta ou indiretamente, de si mesma. "
         "A IA detecta ciclos e avisa.", body_font),
        ("", body_font),
        ("LINKS", section_font),
        ("• Documentação completa: https://github.com/Fabi-Thome/bravo-skills/tree/main/ponto-de-bloqueio", body_font),
        ("• Veja a aba 'Limites do modelo' antes de cadastrar tarefas complexas.", body_font),
    ]

    for i, (text, font) in enumerate(content, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _setup_limites_sheet(wb: Workbook) -> None:
    """Aba com o que NÃO está modelado e por quê."""
    ws = wb.create_sheet("Limites do modelo")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 70

    header_fill = HEADER_FILL
    header_font = HEADER_FONT
    body_font = Font(size=11, color="333333")
    title_font = Font(bold=True, size=14, color="1F2937")

    intro_cell = ws.cell(row=1, column=1, value="Limites conscientes do modelo")
    intro_cell.font = title_font
    ws.merge_cells("A1:B1")

    desc_cell = ws.cell(row=2, column=1,
                        value=("Esta versão do Ponto de Bloqueio é deliberadamente simplificada. "
                               "Veja abaixo o que NÃO é modelado e o que fazer em cada caso."))
    desc_cell.font = body_font
    desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 38

    # Cabeçalho da tabela
    h1 = ws.cell(row=4, column=1, value="O que NÃO é modelado")
    h2 = ws.cell(row=4, column=2, value="Por quê / O que você pode fazer")
    for c in (h1, h2):
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER
    ws.row_dimensions[4].height = 22

    rows = [
        ("Lags variáveis (ex: A termina e B só pode começar 5 dias depois por cura de concreto)",
         "Modele como uma tarefa intermediária 'Cura' entre A e B com duração de 5 dias e contratado fictício 'Aguardo Técnico'."),
        ("Tipos de relação além de Finish-to-Start (FS), como Start-to-Start ou Finish-to-Finish",
         "95% das interfaces de obra são FS. Para os 5% restantes, modele a relação como tarefas distintas com FS entre si."),
        ("Calendários (feriados, fins de semana, intempéries)",
         "O modelo trabalha em dias corridos. Ajuste a duração planejada de cada tarefa considerando os dias úteis reais."),
        ("Recursos (mão-de-obra, equipamentos, equipes compartilhadas)",
         "Análise de recursos é outro problema (RCPSP). Aqui você modela apenas a sequência lógica das interfaces."),
        ("Datas obrigatórias / contratuais (must-start-on, deadline)",
         "Use o campo 'Entrada planejada' para forçar uma data mínima. Para deadlines, monitore manualmente o slip projetado."),
        ("Probabilidades / análise PERT",
         "Modelo determinístico apenas. Para análise de risco probabilístico, use Primavera Risk ou Monte Carlo externo."),
        ("Cronograma interno de cada contratado",
         "Por design — esta skill NÃO precisa do cronograma detalhado. Ela só usa entrada e saída para calcular o caminho crítico das interfaces."),
    ]

    for i, (q, a) in enumerate(rows, start=5):
        cq = ws.cell(row=i, column=1, value=q)
        ca = ws.cell(row=i, column=2, value=a)
        for c in (cq, ca):
            c.font = body_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        ws.row_dimensions[i].height = max(40, 18 * (max(len(q), len(a)) // 60 + 1))


def generate(output_path: str | Path) -> Path:
    """Gera o template e salva no caminho indicado."""
    wb = Workbook()
    ws = wb.active
    _setup_obra_sheet(ws)
    _setup_como_usar_sheet(wb)
    _setup_limites_sheet(wb)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _cli():
    out = Path(__file__).parent.parent / "templates" / "obra-template.xlsx"
    path = generate(out)
    print(f"✅ Template gerado em: {path}")


if __name__ == "__main__":
    _cli()
